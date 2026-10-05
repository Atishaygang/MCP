from fastmcp import FastMCP

from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ============================================================
# MCP SERVER
# ============================================================

mcp = FastMCP(
    "Personal Diary",
    instructions="""
    This server maintains the user's personal diary.

    When the user gives raw, informal or unstructured thoughts,
    transform them into a clean diary entry before calling
    write_diary_entry.

    Preserve the user's actual meaning.

    Never invent events, feelings, people, decisions or facts
    that the user did not express.

    Improve grammar and readability while preserving the user's
    natural voice.
    """
)


# ============================================================
# CONFIGURATION
# ============================================================

# Diary folder will be created beside main.py
DIARY_ROOT = Path(__file__).parent / "diary"

# Indian Standard Time
TIMEZONE = ZoneInfo("Asia/Kolkata")


# ============================================================
# DATE / FILE HELPERS
# ============================================================

def get_current_datetime() -> datetime:
    """
    Get the current date and time in Indian Standard Time.
    """
    return datetime.now(TIMEZONE)


def get_diary_file(date: datetime) -> Path:
    """
    Return the DOCX diary path for a given date.

    Example structure:

    diary/
        2026/
            10/
                2026-10-05.docx
    """

    year = date.strftime("%Y")
    month = date.strftime("%m")

    folder = DIARY_ROOT / year / month

    folder.mkdir(
        parents=True,
        exist_ok=True
    )

    filename = date.strftime("%Y-%m-%d.docx")

    return folder / filename


# ============================================================
# DOCUMENT CREATION
# ============================================================

def create_new_diary_document(date: datetime) -> Document:
    """
    Create a fresh Word diary document for a new day.
    """

    document = Document()

    readable_date = date.strftime("%d %B %Y")
    weekday = date.strftime("%A")

    # --------------------------------------------------------
    # Main title
    # --------------------------------------------------------

    title = document.add_heading(
        "Personal Diary",
        level=0
    )

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --------------------------------------------------------
    # Date
    # --------------------------------------------------------

    date_paragraph = document.add_paragraph()

    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = date_paragraph.add_run(
        f"{weekday}, {readable_date}"
    )

    run.bold = True
    run.font.size = Pt(12)

    document.add_paragraph()

    return document


# ============================================================
# WRITE DIARY
# ============================================================

@mcp.tool()
def write_diary_entry(
    thoughts: str,
    title: str,
    mood: str = "",
    reflection: str = "",
    key_takeaways: list[str] | None = None,
    tags: list[str] | None = None,
) -> str:
    """
    Save a diary entry into today's Word document.

    The user may provide completely unstructured thoughts.

    Before calling this tool, transform those thoughts into
    suitable structured fields.

    Instructions:

    TITLE
    - Create a short meaningful title.
    - Do not exaggerate what happened.

    THOUGHTS
    - Clean the user's grammar and readability.
    - Preserve the user's original meaning and personality.
    - Do not invent information.
    - Do not remove important details.

    MOOD
    - Infer mood only when reasonably evident.
    - Keep it short.
    - Leave empty when mood cannot reasonably be inferred.

    REFLECTION
    - Summarize the deeper meaning or observation.
    - Do not give unnecessary life advice.
    - Do not invent lessons that the user did not imply.

    KEY TAKEAWAYS
    - Important events.
    - Decisions.
    - Lessons.
    - Accomplishments.
    - Things worth remembering.

    TAGS
    - Extract a few useful topics.
    - Examples:
      MCP, College, Gym, Learning, Career, QMTC

    Example:

    User:

    "today was pretty tiring but finally understood MCP.
    made expense tracker and diary server and it actually
    made me understand why MCP is useful"

    Possible call:

    title:
    "MCP Finally Started Making Sense"

    mood:
    "Tired but excited"

    thoughts:
    "Today was tiring, but I finally started understanding
    MCP properly. Building the expense tracker and diary
    server helped me see why MCP can actually be useful."

    reflection:
    "Building practical tools made the concept much easier
    to understand than only studying it theoretically."

    key_takeaways:
    [
        "Understood MCP much better",
        "Built an expense tracker MCP",
        "Built a personal diary MCP"
    ]

    tags:
    [
        "MCP",
        "Learning",
        "Projects"
    ]
    """

    now = get_current_datetime()

    diary_file = get_diary_file(now)

    readable_date = now.strftime("%d %B %Y")
    readable_time = now.strftime("%I:%M %p")

    key_takeaways = key_takeaways or []
    tags = tags or []

    # --------------------------------------------------------
    # Open existing document OR create today's document
    # --------------------------------------------------------

    if diary_file.exists():

        document = Document(diary_file)

    else:

        document = create_new_diary_document(now)

    # --------------------------------------------------------
    # Entry title
    # --------------------------------------------------------

    heading = document.add_heading(
        title,
        level=1
    )

    # --------------------------------------------------------
    # Time
    # --------------------------------------------------------

    time_paragraph = document.add_paragraph()

    time_run = time_paragraph.add_run(
        readable_time
    )

    time_run.italic = True

    # --------------------------------------------------------
    # Mood
    # --------------------------------------------------------

    if mood:

        mood_paragraph = document.add_paragraph()

        mood_label = mood_paragraph.add_run(
            "Mood: "
        )

        mood_label.bold = True

        mood_paragraph.add_run(mood)

    # --------------------------------------------------------
    # Thoughts
    # --------------------------------------------------------

    document.add_heading(
        "Thoughts",
        level=2
    )

    for paragraph_text in thoughts.split("\n"):

        paragraph_text = paragraph_text.strip()

        if paragraph_text:

            document.add_paragraph(
                paragraph_text
            )

    # --------------------------------------------------------
    # Reflection
    # --------------------------------------------------------

    if reflection:

        document.add_heading(
            "Reflection",
            level=2
        )

        document.add_paragraph(
            reflection
        )

    # --------------------------------------------------------
    # Key Takeaways
    # --------------------------------------------------------

    if key_takeaways:

        document.add_heading(
            "Key Takeaways",
            level=2
        )

        for takeaway in key_takeaways:

            document.add_paragraph(
                takeaway,
                style="List Bullet"
            )

    # --------------------------------------------------------
    # Tags
    # --------------------------------------------------------

    if tags:

        document.add_heading(
            "Tags",
            level=2
        )

        formatted_tags = []

        for tag in tags:

            clean_tag = (
                tag.strip()
                .replace(" ", "_")
                .replace("#", "")
            )

            if clean_tag:

                formatted_tags.append(
                    f"#{clean_tag}"
                )

        document.add_paragraph(
            " ".join(formatted_tags)
        )

    # --------------------------------------------------------
    # Separator between entries
    # --------------------------------------------------------

    document.add_paragraph()

    separator = document.add_paragraph(
        "─" * 45
    )

    separator.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph()

    # --------------------------------------------------------
    # Save document
    # --------------------------------------------------------

    document.save(diary_file)

    return (
        "Diary entry saved successfully.\n"
        f"Date: {readable_date}\n"
        f"Time: {readable_time}\n"
        f"Title: {title}\n"
        f"File: {diary_file}"
    )


# ============================================================
# READ DOCX HELPER
# ============================================================

def extract_document_text(
    diary_file: Path
) -> str:
    """
    Read text from a Word diary document.
    """

    document = Document(diary_file)

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            paragraphs.append(text)

    return "\n\n".join(paragraphs)


# ============================================================
# READ TODAY'S DIARY
# ============================================================

@mcp.tool()
def read_today_diary() -> str:
    """
    Read all diary entries written today.
    """

    now = get_current_datetime()

    diary_file = get_diary_file(now)

    if not diary_file.exists():

        return "No diary entries have been written today."

    return extract_document_text(
        diary_file
    )


# ============================================================
# READ DIARY BY DATE
# ============================================================

@mcp.tool()
def read_diary_by_date(
    date: str
) -> str:
    """
    Read the diary from a particular date.

    Required format:

    YYYY-MM-DD

    Example:

    2026-10-05
    """

    try:

        requested_date = datetime.strptime(
            date,
            "%Y-%m-%d"
        )

    except ValueError:

        return (
            "Invalid date format. "
            "Please use YYYY-MM-DD."
        )

    diary_file = get_diary_file(
        requested_date
    )

    if not diary_file.exists():

        return (
            f"No diary entry exists for {date}."
        )

    return extract_document_text(
        diary_file
    )


# ============================================================
# LIST DIARY DATES
# ============================================================

@mcp.tool()
def list_diary_entries() -> list[str]:
    """
    List every date for which a diary document exists.
    """

    if not DIARY_ROOT.exists():

        return []

    files = sorted(
        DIARY_ROOT.rglob("*.docx"),
        reverse=True
    )

    return [
        file.stem
        for file in files
    ]


# ============================================================
# GET DIARY FILE LOCATION
# ============================================================

@mcp.tool()
def get_diary_file_path(
    date: str = ""
) -> str:
    """
    Return the location of a diary Word file.

    If no date is supplied, return today's diary file.

    Optional date format:

    YYYY-MM-DD
    """

    if date:

        try:

            requested_date = datetime.strptime(
                date,
                "%Y-%m-%d"
            )

        except ValueError:

            return (
                "Invalid date format. "
                "Please use YYYY-MM-DD."
            )

    else:

        requested_date = get_current_datetime()

    diary_file = get_diary_file(
        requested_date
    )

    if not diary_file.exists():

        return (
            "No diary exists for that date yet."
        )

    return str(
        diary_file.resolve()
    )

# ============================================================
# MCP RESOURCES
# ============================================================


@mcp.resource("diary://today")
def today_diary_resource() -> str:
    """
    Expose today's diary as an MCP resource.
    """

    now = get_current_datetime()
    diary_file = get_diary_file(now)

    if not diary_file.exists():
        return "No diary entries have been written today."

    return extract_document_text(diary_file)


@mcp.resource("diary://entries")
def diary_entries_resource() -> str:
    """
    Expose the list of available diary dates.
    """

    if not DIARY_ROOT.exists():
        return "No diary entries exist yet."

    files = sorted(
        DIARY_ROOT.rglob("*.docx"),
        reverse=True
    )

    if not files:
        return "No diary entries exist yet."

    dates = [
        file.stem
        for file in files
    ]

    return "\n".join(dates)


@mcp.resource("diary://{date}")
def diary_by_date_resource(date: str) -> str:
    """
    Dynamic resource for reading a specific diary date.

    Example:
        diary://2026-10-05
    """

    try:
        requested_date = datetime.strptime(
            date,
            "%Y-%m-%d"
        )

    except ValueError:
        return "Invalid date. Expected YYYY-MM-DD."

    diary_file = get_diary_file(requested_date)

    if not diary_file.exists():
        return f"No diary exists for {date}."

    return extract_document_text(diary_file)

# ============================================================
# MCP SERVER START
# ============================================================

if __name__ == "__main__":
    mcp.run(
        transport="http",
        host="0.0.0.0",
        port=8000
    )